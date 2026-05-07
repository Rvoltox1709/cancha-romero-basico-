from flask import request, jsonify, render_template, send_file
from app import app, db
from models import Reserva
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io

# --- PÁGINA PRINCIPAL ---
@app.route('/')
def inicio():
    return render_template('index.html')

# --- VER RESERVAS DE UN DÍA ---
@app.route('/reservas', methods=['GET'])
def ver_reservas():
    fecha_str = request.args.get('fecha')
    if fecha_str:
        fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
    else:
        fecha = date.today()
    reservas = Reserva.query.filter_by(fecha=fecha).order_by(Reserva.hora_inicio).all()
    return jsonify([{
        'id': r.id,
        'fecha': str(r.fecha),
        'hora_inicio': r.hora_inicio,
        'hora_fin': r.hora_fin,
        'cliente': r.cliente,
        'telefono': r.telefono,
        'monto_total': r.monto_total,
        'monto_pagado': r.monto_pagado,
        'metodo_pago': r.metodo_pago,
        'pagado': r.pagado
    } for r in reservas])

# --- TODAS LAS RESERVAS ---
@app.route('/reservas/todas', methods=['GET'])
def todas_las_reservas():
    reservas = Reserva.query.order_by(Reserva.fecha, Reserva.hora_inicio).all()
    return jsonify([{
        'id': r.id,
        'fecha': str(r.fecha),
        'hora_inicio': r.hora_inicio,
        'hora_fin': r.hora_fin,
        'cliente': r.cliente,
        'telefono': r.telefono,
        'monto_total': r.monto_total,
        'monto_pagado': r.monto_pagado,
        'metodo_pago': r.metodo_pago,
        'pagado': r.pagado
    } for r in reservas])

# --- AGREGAR RESERVA ---
@app.route('/reservas', methods=['POST'])
def agregar_reserva():
    datos = request.get_json()
    nueva = Reserva(
        fecha        = datetime.strptime(datos['fecha'], '%Y-%m-%d').date(),
        hora_inicio  = datos['hora_inicio'],
        hora_fin     = datos['hora_fin'],
        cliente      = datos['cliente'],
        telefono     = datos.get('telefono', ''),
        monto_total  = datos['monto_total'],
        monto_pagado = datos.get('monto_pagado', 0),
        metodo_pago  = datos.get('metodo_pago', ''),
        pagado       = datos.get('pagado', False)
    )
    db.session.add(nueva)
    db.session.commit()
    return jsonify({'mensaje': 'Reserva creada', 'id': nueva.id}), 201

# --- EDITAR RESERVA ---
@app.route('/reservas/<int:id>', methods=['PUT'])
def editar_reserva(id):
    reserva = Reserva.query.get_or_404(id)
    datos = request.get_json()
    reserva.cliente      = datos.get('cliente', reserva.cliente)
    reserva.telefono     = datos.get('telefono', reserva.telefono)
    reserva.hora_inicio  = datos.get('hora_inicio', reserva.hora_inicio)
    reserva.hora_fin     = datos.get('hora_fin', reserva.hora_fin)
    reserva.monto_total  = datos.get('monto_total', reserva.monto_total)
    reserva.monto_pagado = datos.get('monto_pagado', reserva.monto_pagado)
    reserva.metodo_pago  = datos.get('metodo_pago', reserva.metodo_pago)
    reserva.pagado       = datos.get('pagado', reserva.pagado)
    db.session.commit()
    return jsonify({'mensaje': 'Reserva actualizada'})

# --- ELIMINAR RESERVA ---
@app.route('/reservas/<int:id>', methods=['DELETE'])
def eliminar_reserva(id):
    reserva = Reserva.query.get_or_404(id)
    db.session.delete(reserva)
    db.session.commit()
    return jsonify({'mensaje': 'Reserva eliminada'})

# --- RESUMEN DEL DÍA ---
@app.route('/resumen/dia', methods=['GET'])
def resumen_dia():
    fecha_str = request.args.get('fecha')
    if fecha_str:
        fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
    else:
        fecha = date.today()
    reservas = Reserva.query.filter_by(fecha=fecha).all()
    total_esperado  = sum(r.monto_total for r in reservas)
    total_cobrado   = sum(r.monto_pagado for r in reservas)
    return jsonify({
        'fecha': str(fecha),
        'total_reservas': len(reservas),
        'total_esperado': total_esperado,
        'total_cobrado': total_cobrado,
        'total_pendiente': total_esperado - total_cobrado
    })

# --- RESUMEN DEL MES ---
@app.route('/resumen/mes', methods=['GET'])
def resumen_mes():
    anio = int(request.args.get('anio', date.today().year))
    mes  = int(request.args.get('mes', date.today().month))
    reservas = Reserva.query.filter(
        db.extract('year', Reserva.fecha) == anio,
        db.extract('month', Reserva.fecha) == mes
    ).all()
    total_esperado  = sum(r.monto_total for r in reservas)
    total_cobrado   = sum(r.monto_pagado for r in reservas)
    return jsonify({
        'anio': anio,
        'mes': mes,
        'total_reservas': len(reservas),
        'total_esperado': total_esperado,
        'total_cobrado': total_cobrado,
        'total_pendiente': total_esperado - total_cobrado
    })

# --- EXPORTAR EXCEL ---
@app.route('/exportar/excel', methods=['GET'])
def exportar_excel():
    mes  = int(request.args.get('mes', date.today().month))
    anio = int(request.args.get('anio', date.today().year))
    reservas = Reserva.query.filter(
        db.extract('year', Reserva.fecha) == anio,
        db.extract('month', Reserva.fecha) == mes
    ).order_by(Reserva.fecha, Reserva.hora_inicio).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reservas {mes}-{anio}"

    ws.merge_cells('A1:J1')
    ws['A1'] = f'Reservas Cancha - {mes}/{anio}'
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='1a3a1a')
    ws['A1'].alignment = Alignment(horizontal='center')

    cabeceras = ['Fecha','Hora inicio','Hora fin','Cliente','Teléfono','Monto total','Monto pagado','Pendiente','Método pago','Estado']
    for col, cab in enumerate(cabeceras, 1):
        celda = ws.cell(row=2, column=col, value=cab)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='2d6a2d')
        celda.alignment = Alignment(horizontal='center')

    for fila, r in enumerate(reservas, 3):
        pendiente = r.monto_total - r.monto_pagado
        ws.cell(row=fila, column=1, value=str(r.fecha))
        ws.cell(row=fila, column=2, value=r.hora_inicio)
        ws.cell(row=fila, column=3, value=r.hora_fin)
        ws.cell(row=fila, column=4, value=r.cliente)
        ws.cell(row=fila, column=5, value=r.telefono)
        ws.cell(row=fila, column=6, value=r.monto_total)
        ws.cell(row=fila, column=7, value=r.monto_pagado)
        ws.cell(row=fila, column=8, value=pendiente)
        ws.cell(row=fila, column=9, value=r.metodo_pago or '')
        ws.cell(row=fila, column=10, value='Pagado' if r.pagado else 'Pendiente')
        color = 'e8f5e1' if r.pagado else 'fff3cd'
        for col in range(1, 11):
            ws.cell(row=fila, column=col).fill = PatternFill('solid', fgColor=color)

    fila_total = len(reservas) + 3
    ws.cell(row=fila_total, column=3, value='TOTAL').font = Font(bold=True)
    ws.cell(row=fila_total, column=6, value=sum(r.monto_total for r in reservas)).font = Font(bold=True)
    ws.cell(row=fila_total, column=7, value=sum(r.monto_pagado for r in reservas)).font = Font(bold=True)
    ws.cell(row=fila_total, column=8, value=sum(r.monto_total - r.monto_pagado for r in reservas)).font = Font(bold=True)

    anchos = [12,12,10,20,15,13,14,12,15,12]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'cancha_reservas_{mes}_{anio}.xlsx')